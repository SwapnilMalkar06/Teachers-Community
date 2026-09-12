import os
import win32com.client
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right']:
        if border_name in kwargs:
            edge = OxmlElement(f'w:{border_name}')
            for key, val in kwargs[border_name].items():
                edge.set(qn(f'w:{key}'), str(val))
            tcBorders.append(edge)
    tcPr.append(tcBorders)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_synopsis_docx(output_docx_path):
    doc = Document()
    
    # Page setup - Margins (0.55 inch top/bottom, 0.6 inch left/right to ensure clean 2-page fit)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)
        
    # Default Font: Times New Roman
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    
    # Header Paragraphs
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(2)
    p1.paragraph_format.line_spacing = 1.1
    run1 = p1.add_run("Kolhapur Institute of Technology's\n")
    run1.font.bold = True
    run1.font.size = Pt(13)
    run2 = p1.add_run("Institute of Management Education and Research (Autonomous), Kolhapur.")
    run2.font.bold = True
    run2.font.size = Pt(13)
    
    # Horizontal line rule
    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_before = Pt(1)
    p_line.paragraph_format.space_after = Pt(3)
    run_line = p_line.add_run("------------------------------------------------------------------------------------------------------------------------")
    run_line.font.size = Pt(9.5)
    run_line.font.color.rgb = RGBColor(80, 80, 80)
    
    # MCA Info
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(10)
    p2.paragraph_format.line_spacing = 1.15
    r_mca1 = p2.add_run("MCA Department.\n")
    r_mca1.font.bold = True
    r_mca1.font.size = Pt(11.5)
    r_mca2 = p2.add_run("MCA-II (Sem-III)-2026-27.")
    r_mca2.font.bold = True
    r_mca2.font.size = Pt(11.5)
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(12)
    r_title = p_title.add_run("Mini-Project Synopsis")
    r_title.font.bold = True
    r_title.font.size = Pt(15)
    
    # Date (Right aligned)
    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_date.paragraph_format.space_before = Pt(0)
    p_date.paragraph_format.space_after = Pt(6)
    r_date = p_date.add_run("Date: 03 / 09 / 2026")
    r_date.font.bold = True
    r_date.font.size = Pt(11)
    
    # Table Setup - 8 rows, 2 columns
    table = doc.add_table(rows=8, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    col_widths = [Inches(4.6), Inches(2.7)]
    
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = col_widths[i]
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            set_cell_border(cell, 
                            top=dict(sz=6, val='single', color='000000'),
                            bottom=dict(sz=6, val='single', color='000000'),
                            left=dict(sz=6, val='single', color='000000'),
                            right=dict(sz=6, val='single', color='000000'))
            
    # Row 0: Roll Number & Division
    cell_r0_c0 = table.cell(0, 0)
    cell_r0_c1 = table.cell(0, 1)
    
    p = cell_r0_c0.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Roll Number:- ")
    r.bold = True
    p.add_run("MCA2-2026-042")
    
    p = cell_r0_c1.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Division:- ")
    r.bold = True
    p.add_run("A")
    
    # Row 1: Name of Student (Merged)
    cell_r1 = table.cell(1, 0)
    cell_r1.merge(table.cell(1, 1))
    p = cell_r1.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Name of Student: ")
    r.bold = True
    p.add_run("Ms. Ashwini Sawant")
    
    # Row 2: Partner Roll Number & Division
    cell_r2_c0 = table.cell(2, 0)
    cell_r2_c1 = table.cell(2, 1)
    
    p = cell_r2_c0.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Partner Roll Number( Group Project): ")
    r.bold = True
    p.add_run("N/A (Individual Project)")
    
    p = cell_r2_c1.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Division: ")
    r.bold = True
    p.add_run("N/A")
    
    # Row 3: Name of Partner (Merged)
    cell_r3 = table.cell(3, 0)
    cell_r3.merge(table.cell(3, 1))
    p = cell_r3.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Name of Partner(Group Project): ")
    r.bold = True
    p.add_run("N/A")
    
    # Row 4: Project Title (Merged)
    cell_r4 = table.cell(4, 0)
    cell_r4.merge(table.cell(4, 1))
    p = cell_r4.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Project Title: ")
    r.bold = True
    p.add_run("Web-based Faculty Portfolio & Academic Content Management System (CMS)")
    
    # Row 5: Name of Organization & Address (Merged)
    cell_r5 = table.cell(5, 0)
    cell_r5.merge(table.cell(5, 1))
    p = cell_r5.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Name of Organization & Address (For Live Project):- ")
    r.bold = True
    p.add_run("In-House Academic Project, KIT's Institute of Management Education and Research (Autonomous), Kolhapur.")
    
    # Row 6: Project Guide (Merged)
    cell_r6 = table.cell(6, 0)
    cell_r6.merge(table.cell(6, 1))
    p = cell_r6.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Project Guide: ")
    r.bold = True
    p.add_run("Prof. Dr. A. B. Patil / MCA Department Faculty Guide")
    
    # Row 7: Project Description (Merged)
    cell_r7 = table.cell(7, 0)
    cell_r7.merge(table.cell(7, 1))
    
    p0 = cell_r7.paragraphs[0]
    p0.paragraph_format.space_before = Pt(4)
    p0.paragraph_format.space_after = Pt(4)
    r_desc = p0.add_run("Project Description:")
    r_desc.bold = True
    r_desc.font.size = Pt(11.5)
    
    def add_subsection(title_text, content_paragraphs):
        p_sub = cell_r7.add_paragraph()
        p_sub.paragraph_format.space_before = Pt(5)
        p_sub.paragraph_format.space_after = Pt(2)
        p_sub.paragraph_format.line_spacing = 1.1
        r_sub = p_sub.add_run(title_text)
        r_sub.bold = True
        r_sub.font.size = Pt(10.5)
        
        for item in content_paragraphs:
            p_c = cell_r7.add_paragraph()
            p_c.paragraph_format.space_before = Pt(0)
            p_c.paragraph_format.space_after = Pt(2)
            p_c.paragraph_format.line_spacing = 1.1
            if isinstance(item, tuple):
                label, val = item
                r_lbl = p_c.add_run(label)
                r_lbl.bold = True
                p_c.add_run(val)
            else:
                p_c.add_run(item)

    # a. Introduction:
    add_subsection("a. Introduction:", [
        "The Web-based Faculty Portfolio & Academic Content Management System (CMS) is a modern, full-stack web platform designed specifically for higher education institutions to bridge the gap between academic resource distribution, research dissemination, and student engagement.",
        "The platform serves as a unified digital portfolio and CMS for Prof. Ashwini Sawant, enabling students to access course materials, downloadable lecture notes, slide decks, video tutorials, question banks, and published research, while providing faculty with a secure Admin Dashboard to manage and update content dynamically."
    ])
    
    # b. Objectives:
    add_subsection("b. Objectives:", [
        ("- Centralized Digital Profile: ", "To construct a unified online portfolio showcasing academic credentials, teaching experience, research interests, and committee roles."),
        ("- Resource Distribution Portal: ", "To deliver a structured Teaching Portal where students can search, filter, and download subject-wise notes, PPTs, video links, and question banks."),
        ("- Research & Publication Repository: ", "To systematically catalog PhD research, peer-reviewed journal papers, and conference proceedings with DOI links and downloadable PDFs."),
        ("- Student Engagement Analytics: ", "To implement real-time engagement and view duration tracking on academic blogs using browser visibility APIs."),
        ("- Secure Content Management (CMS): ", "To equip faculty with an intuitive Admin Dashboard for complete CRUD management of profile content, resources, and contact inquiries.")
    ])
    
    # c. Problem Statement:
    add_subsection("c. Problem Statement:", [
        "In traditional academic setups, educational resources and research records are distributed across fragmented channels (social media groups, cloud drives, email attachments, physical notice boards). This creates key operational challenges:",
        "1. Inconvenient Student Access: Students frequently miss critical lecture notes, model answers, and academic updates.",
        "2. Unstructured Research Record: Faculty publications, conference talks, and certifications are poorly indexed and difficult for external researchers to discover.",
        "3. Technical Overhead: Traditional websites require manual developer intervention to publish new updates or documents.",
        "4. Lack of Engagement Insight: Faculty members cannot measure student interaction or reading metrics for shared academic content."
    ])
    
    # d. Proposed Solution:
    add_subsection("d. Proposed Solution:", [
        "The proposed system delivers a full-stack Next.js 14 web platform integrated with Prisma ORM and MySQL/SQLite database. It comprises two main operational modules:",
        "1. Public Student & Peer Portal: A responsive user interface containing Home Profile, Teaching Resources (Notes, PPTs, Videos, Question Banks), Research Publications Hub, FDP/Workshops timeline, Certificate lightbox gallery, Analytics Blog, and Interactive Contact form.",
        "2. Admin CMS Dashboard: A secure, authenticated control center enabling full CRUD control over profile details, academic subjects, resources, research records, certificates, blog posts, and contact inbox messages."
    ])
    
    # e. Front, Back End and Database :
    add_subsection("e. Front, Back End and Database :", [
        ("- Front End: ", "Next.js 14 (App Router), React 18, TypeScript, Tailwind CSS, Lucide React Icons, Framer Motion."),
        ("- Back End: ", "Node.js environment, Next.js API Routes (RESTful API), Server Actions, Prisma Client ORM."),
        ("- Database: ", "MySQL Database / SQLite (configured via Prisma ORM with relational schema integrity).")
    ])
    
    # f. System Requirements(H/w-S/w):
    add_subsection("f. System Requirements(H/w-S/w):", [
        ("1. Hardware Requirements:", ""),
        ("   - Processor: ", "Intel Core i3 / i5 / AMD Ryzen 5 or higher"),
        ("   - System Memory (RAM): ", "8 GB minimum"),
        ("   - Storage Space: ", "20 GB available SSD space"),
        ("   - Display Resolution: ", "1920 x 1080 (Full HD)"),
        ("   - Internet Connection: ", "High-speed broadband network connection"),
        ("2. Software Requirements:", ""),
        ("   - Operating System: ", "Windows 10 / 11 (64-bit), Linux, or macOS"),
        ("   - Environment Runtime: ", "Node.js v18.x or v20.x LTS"),
        ("   - Database Engine: ", "MySQL 8.0+ / SQLite 3"),
        ("   - Core Framework: ", "Next.js 14 (React 18, TypeScript 5)"),
        ("   - Package Manager: ", "npm / yarn / pnpm"),
        ("   - Web Browser: ", "Google Chrome, Mozilla Firefox, Microsoft Edge, Safari")
    ])
    
    doc.save(output_docx_path)
    print(f"DOCX created successfully at: {output_docx_path}")

def convert_docx_to_pdf(docx_path, pdf_path):
    abs_docx = os.path.abspath(docx_path)
    abs_pdf = os.path.abspath(pdf_path)
    
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    try:
        doc = word.Documents.Open(abs_docx)
        # 17 represents wdFormatPDF
        doc.SaveAs(abs_pdf, FileFormat=17)
        doc.Close()
        print(f"PDF created successfully at: {abs_pdf}")
    except Exception as e:
        print("Error during DOCX to PDF conversion:", e)
    finally:
        word.Quit()

if __name__ == "__main__":
    docx_file = r"d:\Ashwini_Savant\Mini_Project_Synopsis.docx"
    pdf_file = r"d:\Ashwini_Savant\Mini_Project_Synopsis.pdf"
    
    create_synopsis_docx(docx_file)
    convert_docx_to_pdf(docx_file, pdf_file)
